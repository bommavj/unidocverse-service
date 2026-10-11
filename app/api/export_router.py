"""
export_router.py
================
Universal export endpoints — returns XLSX or CSV for any domain or document list.
"""
import io
import csv
import logging
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import text

from app.core.database import get_db
from app.core.tenant_scope import scope_sql, tenant_params, current_tenant

router = APIRouter(prefix="/api/export", tags=["export"])
logger = logging.getLogger(__name__)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _xlsx_response(filename: str, sheets: dict) -> StreamingResponse:
    """sheets = {sheet_name: (headers, rows)}"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    HEADER_FILL  = PatternFill("solid", fgColor="4F46E5")
    HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
    HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

    for sheet_name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.fill  = HEADER_FILL
            cell.font  = HEADER_FONT
            cell.alignment = HEADER_ALIGN

        for row in rows:
            ws.append([str(v) if v is not None else "" for v in row])

        # Auto-width
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def _csv_response(filename: str, headers: list, rows: list) -> StreamingResponse:
    buf = io.StringIO()
    w   = csv.writer(buf)
    w.writerow(headers)
    w.writerows([[str(v) if v is not None else "" for v in r] for r in rows])
    buf.seek(0)
    return StreamingResponse(
        io.BytesIO(buf.getvalue().encode()),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def _ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")


# ── Documents ─────────────────────────────────────────────────────────────────

def scoped(sql: str, params=None):
    """Tenant-scoped query: domain tables are filtered to the current tenant's documents."""
    return text(scope_sql(sql)), tenant_params(params)


@router.get("/documents")
def export_documents(
    fmt: str = Query("xlsx", regex="^(xlsx|csv)$"),
    doc_type: Optional[str] = None,
    folder: Optional[str] = None,
    db: Session = Depends(get_db)
):
    q = ("SELECT filename, doc_type, NULL AS folder_name, file_size, status, created_at FROM documents "
         "WHERE tenant_id = :tenant_id")         # documents has no folder column
    p: dict = {"tenant_id": current_tenant()}
    if doc_type: q += " AND doc_type=:dt"; p["dt"] = doc_type
    # `folder` is accepted for compatibility; documents has no folder column to filter on
    q += " ORDER BY created_at DESC"
    rows = db.execute(text(q), p).fetchall()

    headers = ["Filename", "Type", "Folder", "Size (KB)", "Status", "Uploaded At"]
    data    = [(r.filename, r.doc_type, r.folder_name,
                round((r.file_size or 0) / 1024, 1), r.status,
                str(r.created_at)[:19]) for r in rows]

    fn = f"documents_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, headers, data)
    return _xlsx_response(fn, {"Documents": (headers, data)})


# ── Legal ─────────────────────────────────────────────────────────────────────

@router.get("/legal")
def export_legal(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    contracts = db.execute(*scoped("""
        SELECT lc.contract_name, lc.counterparty, lc.notice_period,
               lc.expiry_date, lc.risk_rating, lc.status,
               COALESCE(d.doc_type,'contract') AS doc_type
        FROM legal_contracts lc
        LEFT JOIN documents d ON d.id = lc.document_id
        ORDER BY lc.created_at DESC
    """)).fetchall()
    clauses = db.execute(*scoped(
        "SELECT clause_type, clause_name, clause_text, is_deviation, governing_law FROM legal_clauses ORDER BY created_at DESC"
    )).fetchall()

    ch = ["Contract Name", "Counterparty", "Notice Period", "Expiry Date", "Risk", "Status", "Type"]
    cd = [(r.contract_name, r.counterparty, r.notice_period, str(r.expiry_date), r.risk_rating, r.status, r.doc_type) for r in contracts]
    clh = ["Clause Type", "Name", "Text", "Is Deviation", "Governing Law"]
    cld = [(r.clause_type, r.clause_name, (r.clause_text or "")[:200], r.is_deviation, r.governing_law) for r in clauses]

    fn = f"legal_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, ch, cd)
    return _xlsx_response(fn, {"Contracts": (ch, cd), "Clauses": (clh, cld)})


# ── Healthcare ────────────────────────────────────────────────────────────────

@router.get("/healthcare")
def export_healthcare(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    patients = db.execute(*scoped(
        "SELECT patient_name, patient_dob, chart_score, consent_signed, insurance_verified FROM patient_charts ORDER BY created_at DESC"
    )).fetchall()
    biomarkers = db.execute(*scoped(
        "SELECT biomarker_name, biomarker_value, unit, reference_range, is_abnormal, physician FROM patient_biomarkers ORDER BY created_at DESC"
    )).fetchall()

    ph  = ["Patient Name", "DOB", "Chart Score", "Consent Signed", "Insurance Verified"]
    pd_ = [(r.patient_name, str(r.patient_dob), r.chart_score, r.consent_signed, r.insurance_verified) for r in patients]
    bh  = ["Biomarker", "Value", "Unit", "Reference Range", "Abnormal", "Physician"]
    bd  = [(r.biomarker_name, r.biomarker_value, r.unit, r.reference_range, r.is_abnormal, r.physician) for r in biomarkers]

    fn = f"healthcare_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, ph, pd_)
    return _xlsx_response(fn, {"Patients": (ph, pd_), "Biomarkers": (bh, bd)})


# ── Finance ───────────────────────────────────────────────────────────────────

@router.get("/finance")
def export_finance(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    matches = db.execute(*scoped(
        "SELECT invoice_id, vendor_name, invoice_val, po_val, receipt_val, match_status, is_approved FROM finance_matches ORDER BY created_at DESC"
    )).fetchall()
    anomalies = db.execute(*scoped(
        "SELECT anomaly_type, title, ref_code, amount FROM finance_anomalies ORDER BY created_at DESC"
    )).fetchall()

    mh = ["Invoice ID", "Vendor", "Invoice Value", "PO Value", "Receipt Value", "Match Status", "Approved"]
    md = [(r.invoice_id, r.vendor_name, float(r.invoice_val or 0), float(r.po_val or 0),
           float(r.receipt_val or 0), r.match_status, r.is_approved) for r in matches]
    ah = ["Anomaly Type", "Title", "Reference", "Amount"]
    ad = [(r.anomaly_type, r.title, r.ref_code, float(r.amount or 0)) for r in anomalies]

    fn = f"finance_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, mh, md)
    return _xlsx_response(fn, {"Invoice Matches": (mh, md), "Anomalies": (ah, ad)})


# ── Real Estate ───────────────────────────────────────────────────────────────

@router.get("/real-estate")
def export_real_estate(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    props = db.execute(*scoped(
        "SELECT address, buyer_name, price, progress, has_appraisal, has_clear_title, is_escrow_funded FROM real_estate_properties ORDER BY created_at DESC"
    )).fetchall()
    milestones = db.execute(*scoped("""
        SELECT p.address, m.name, m.due_date, m.is_completed
        FROM real_estate_milestones m
        JOIN real_estate_properties p ON p.id = m.property_id
        ORDER BY m.due_date
    """)).fetchall()

    prh = ["Address", "Buyer", "Price", "Progress %", "Appraisal", "Clear Title", "Escrow Funded"]
    prd = [(r.address, r.buyer_name, float(r.price or 0), r.progress, r.has_appraisal, r.has_clear_title, r.is_escrow_funded) for r in props]
    mlh = ["Address", "Milestone", "Due Date", "Completed"]
    mld = [(r.address, r.name, str(r.due_date), r.is_completed) for r in milestones]

    fn = f"real_estate_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, prh, prd)
    return _xlsx_response(fn, {"Properties": (prh, prd), "Milestones": (mlh, mld)})


# ── HR ────────────────────────────────────────────────────────────────────────

@router.get("/hr")
def export_hr(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    candidates = db.execute(*scoped(
        "SELECT candidate_name, role, experience_years, match_score, skills, feedback FROM hr_candidates ORDER BY created_at DESC"
    )).fetchall()
    onboarding = db.execute(*scoped(
        "SELECT employee_name, title, due_date, is_done FROM hr_onboarding_tasks ORDER BY due_date"
    )).fetchall()

    ch  = ["Name", "Role", "Experience (yrs)", "Match Score", "Skills", "Feedback"]
    cd  = [(r.candidate_name, r.role, r.experience_years, r.match_score, ", ".join(r.skills or []), r.feedback) for r in candidates]
    oh  = ["Employee", "Task", "Due Date", "Done"]
    od  = [(r.employee_name, r.title, str(r.due_date), r.is_done) for r in onboarding]

    fn = f"hr_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, ch, cd)
    return _xlsx_response(fn, {"Candidates": (ch, cd), "Onboarding Tasks": (oh, od)})


# ── Supply Chain ──────────────────────────────────────────────────────────────

@router.get("/supply-chain")
def export_supply_chain(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    shipments = db.execute(*scoped(
        "SELECT bol_number, carrier_name, origin, destination, customs_status, is_reconciled FROM supply_chain_shipments ORDER BY created_at DESC"
    )).fetchall()
    suppliers = db.execute(*scoped(
        "SELECT name, type, expiry_date, sanctions_status FROM supply_chain_suppliers ORDER BY created_at DESC"
    )).fetchall()

    sh  = ["BOL Number", "Carrier", "Origin", "Destination", "Customs", "Reconciled"]
    sd  = [(r.bol_number, r.carrier_name, r.origin, r.destination, r.customs_status, r.is_reconciled) for r in shipments]
    suh = ["Supplier", "Type", "Expiry Date", "Sanctions Status"]
    sud = [(r.name, r.type, str(r.expiry_date), r.sanctions_status) for r in suppliers]

    fn = f"supply_chain_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, sh, sd)
    return _xlsx_response(fn, {"Shipments": (sh, sd), "Suppliers": (suh, sud)})


# ── Agency (Clients, Policies) ────────────────────────────────────────────────

@router.get("/agency")
def export_agency(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    t = {"tenant_id": current_tenant()}
    clients = db.execute(text(
        "SELECT name, email, phone, status, city, state FROM clients WHERE tenant_id = :tenant_id "
        "ORDER BY created_at DESC LIMIT 1000"), t).fetchall()
    policies = db.execute(text(
        "SELECT policy_number, line_of_business, carrier_name, premium_amount, effective_date, expiry_date, status "
        "FROM policies WHERE tenant_id = :tenant_id ORDER BY created_at DESC LIMIT 1000"), t).fetchall()

    clh = ["Name", "Email", "Phone", "Status", "City", "State"]
    cld = [(r.name, r.email, r.phone, r.status, r.city, r.state) for r in clients]
    plh = ["Policy #", "Line of Business", "Carrier", "Premium", "Effective", "Expiry", "Status"]
    pld = [(r.policy_number, r.line_of_business, r.carrier_name,
            float(r.premium_amount or 0), str(r.effective_date),
            str(r.expiry_date), r.status) for r in policies]

    fn = f"agency_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, clh, cld)
    return _xlsx_response(fn, {"Clients": (clh, cld), "Policies": (plh, pld)})


# ── Insurance Automation Log ──────────────────────────────────────────────────

@router.get("/automation-log")
def export_automation_log(fmt: str = Query("xlsx", regex="^(xlsx|csv)$"), db: Session = Depends(get_db)):
    rows = db.execute(*scoped(
        "SELECT rule_name, entity_type, entity_id, action_taken, result, detail, error, triggered_at "
        "FROM insurance_automation_log ORDER BY triggered_at DESC LIMIT 5000")).fetchall()

    h = ["Rule", "Entity Type", "Entity ID", "Action", "Result", "Detail", "Error", "Triggered At"]
    d = [(r.rule_name, r.entity_type, r.entity_id, r.action_taken, r.result,
          (r.detail or "")[:200], (r.error or "")[:200], str(r.triggered_at)[:19]) for r in rows]

    fn = f"automation_log_{_ts()}.{fmt}"
    if fmt == "csv":
        return _csv_response(fn, h, d)
    return _xlsx_response(fn, {"Automation Log": (h, d)})
