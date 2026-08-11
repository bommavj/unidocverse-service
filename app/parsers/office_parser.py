"""Office document parser (Excel, Word, PowerPoint)"""

import logging
from typing import Dict, Any
from pathlib import Path

logger = logging.getLogger(__name__)


class OfficeParser:
    """Parse Microsoft Office documents"""

    def parse(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse Office documents"""
        ext = Path(filename).suffix.lower()

        if ext in ['.xlsx', '.xls', '.xlsm']:
            return self._parse_excel(file_path)
        elif ext in ['.docx', '.doc']:
            return self._parse_word(file_path)
        elif ext in ['.pptx', '.ppt']:
            return self._parse_powerpoint(file_path)

        return {'raw_text': '', 'error': 'Unsupported Office format'}

    def _parse_excel(self, file_path: str) -> Dict[str, Any]:
        """Parse Excel files"""
        try:
            import pandas as pd

            excel_file = pd.ExcelFile(file_path)

            result = {
                'raw_text': '',
                'tables': [],
                'metadata': {'sheets': excel_file.sheet_names},
                'parser_used': 'pandas',
                'confidence': 0.95
            }

            text_parts = []

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)

                # Convert to text
                sheet_text = f"Sheet: {sheet_name}\n"
                sheet_text += f"Columns: {', '.join(str(col) for col in df.columns)}\n\n"
                sheet_text += df.to_string()
                text_parts.append(sheet_text)

                # Store as table
                result['tables'].append({
                    'sheet': sheet_name,
                    'data': df.values.tolist(),
                    'columns': df.columns.tolist(),
                    'rows': len(df),
                    'cols': len(df.columns)
                })

            result['raw_text'] = "\n\n".join(text_parts)
            logger.info(f"✅ Excel parsed: {len(text_parts)} sheets")
            return result

        except Exception as e:
            logger.error(f"Excel parsing failed: {e}")
            return {'raw_text': '', 'error': str(e), 'confidence': 0.0}

    def _parse_word(self, file_path: str) -> Dict[str, Any]:
        """Parse Word documents"""
        try:
            import docx

            doc = docx.Document(file_path)
            text_parts = [para.text for para in doc.paragraphs if para.text.strip()]

            return {
                'raw_text': "\n\n".join(text_parts),
                'tables': [],
                'metadata': {'paragraphs': len(text_parts)},
                'parser_used': 'python-docx',
                'confidence': 0.9
            }

        except Exception as e:
            logger.error(f"Word parsing failed: {e}")
            return {'raw_text': '', 'error': str(e)}

    def _parse_powerpoint(self, file_path: str) -> Dict[str, Any]:
        """Parse PowerPoint files"""
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            text_parts = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = f"Slide {slide_num}:\n"
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        slide_text += shape.text + "\n"
                text_parts.append(slide_text)

            return {
                'raw_text': "\n\n".join(text_parts),
                'tables': [],
                'metadata': {'slides': len(prs.slides)},
                'parser_used': 'python-pptx',
                'confidence': 0.85
            }

        except Exception as e:
            logger.error(f"PowerPoint parsing failed: {e}")
            return {'raw_text': '', 'error': str(e)}
        
        
# """Office document parser (Excel, Word, PowerPoint)"""
# 
# import logging
# from typing import Dict, Any
# from pathlib import Path
# 
# logger = logging.getLogger(__name__)
# 
# 
# class OfficeParser:
#     """Parse Microsoft Office documents"""
# 
#     def parse(self, file_path: str, filename: str) -> Dict[str, Any]:
#         """Parse Office documents"""
#         ext = Path(filename).suffix.lower()
# 
#         if ext in ['.xlsx', '.xls', '.xlsm']:
#             return self._parse_excel(file_path)
#         elif ext in ['.docx', '.doc']:
#             return self._parse_word(file_path)
#         elif ext in ['.pptx', '.ppt']:
#             return self._parse_powerpoint(file_path)
# 
#         return {'raw_text': '', 'error': 'Unsupported Office format'}
# 
#     def _parse_excel(self, file_path: str) -> Dict[str, Any]:
#         """Parse Excel files"""
#         import pandas as pd
# 
#         try:
#             excel_file = pd.ExcelFile(file_path)
# 
#             result = {
#                 'raw_text': '',
#                 'tables': [],
#                 'metadata': {'sheets': excel_file.sheet_names},
#                 'parser_used': 'pandas',
#                 'confidence': 0.95
#             }
# 
#             text_parts = []
# 
#             for sheet_name in excel_file.sheet_names:
#                 df = pd.read_excel(file_path, sheet_name=sheet_name)
# 
#                 # Convert to text
#                 sheet_text = f"Sheet: {sheet_name}\n"
#                 sheet_text += f"Columns: {', '.join(str(col) for col in df.columns)}\n\n"
#                 sheet_text += df.to_string()
#                 text_parts.append(sheet_text)
# 
#                 # Store as table
#                 result['tables'].append({
#                     'sheet': sheet_name,
#                     'data': df.values.tolist(),
#                     'columns': df.columns.tolist(),
#                     'rows': len(df),
#                     'cols': len(df.columns)
#                 })
# 
#             result['raw_text'] = "\n\n".join(text_parts)
#             return result
# 
#         except Exception as e:
#             logger.error(f"Excel parsing failed: {e}")
#             return {'raw_text': '', 'error': str(e)}
# 
#     def _parse_word(self, file_path: str) -> Dict[str, Any]:
#         """Parse Word documents"""
#         try:
#             import docx
# 
#             doc = docx.Document(file_path)
# 
#             # Extract paragraphs
#             text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
# 
#             # Extract tables
#             tables = []
#             for table in doc.tables:
#                 table_data = []
#                 for row in table.rows:
#                     row_data = [cell.text for cell in row.cells]
#                     table_data.append(row_data)
#                 tables.append({
#                     'data': table_data,
#                     'rows': len(table_data),
#                     'cols': len(table_data[0]) if table_data else 0
#                 })
# 
#             return {
#                 'raw_text': "\n\n".join(text_parts),
#                 'tables': tables,
#                 'metadata': {'paragraphs': len(text_parts)},
#                 'parser_used': 'python-docx',
#                 'confidence': 0.9
#             }
# 
#         except Exception as e:
#             logger.error(f"Word parsing failed: {e}")
#             return {'raw_text': '', 'error': str(e)}
# 
#     def _parse_powerpoint(self, file_path: str) -> Dict[str, Any]:
#         """Parse PowerPoint files"""
#         try:
#             from pptx import Presentation
# 
#             prs = Presentation(file_path)
# 
#             text_parts = []
#             for slide_num, slide in enumerate(prs.slides, 1):
#                 slide_text = f"Slide {slide_num}:\n"
#                 for shape in slide.shapes:
#                     if hasattr(shape, 'text'):
#                         slide_text += shape.text + "\n"
#                 text_parts.append(slide_text)
# 
#             return {
#                 'raw_text': "\n\n".join(text_parts),
#                 'tables': [],
#                 'metadata': {'slides': len(prs.slides)},
#                 'parser_used': 'python-pptx',
#                 'confidence': 0.85
#             }
# 
#         except Exception as e:
#             logger.error(f"PowerPoint parsing failed: {e}")
#             return {'raw_text': '', 'error': str(e)}